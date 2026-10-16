"""Presentazione tutorial (PowerPoint) dal JSON finale, per un formatore che
la proietta in aula.

Layout neutro 16:9 su sfondo bianco, colori dell'HTML. Sulla slide va ciò che
si proietta (titolo, immagine o testo essenziale); il testo completo e la
fonte di ogni step vanno nelle note del relatore. Nessun tetto alla
lunghezza: il testo troppo lungo per la slide si tronca con "…" e resta
intero nelle note.
"""

from __future__ import annotations

import io
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from doc_parser.output.strings import get_strings

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
MARGIN = Inches(0.5)
CONTENT_TOP = Inches(1.55)
IMAGE_WIDTH = Inches(12.3)
IMAGE_HEIGHT = Inches(5.6)

TEXT = RGBColor(0x21, 0x25, 0x29)
ACCENT = RGBColor(0x1A, 0x73, 0xE8)
SECONDARY = RGBColor(0x5F, 0x63, 0x68)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

AGENDA_PER_SLIDE = 10
GLOSSARY_PER_SLIDE = 8
NOTES_PER_SLIDE = 6
_BLANK_LAYOUT = 6


def _truncate(text: str, limit: int) -> str:
    text = " ".join(str(text or "").split())
    return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"


def _chunks(items: list, size: int) -> list[list]:
    return [items[i : i + size] for i in range(0, len(items), size)]


def _paged(title: str, page: int, pages: int) -> str:
    return f"{title} ({page}/{pages})" if pages > 1 else title


class _Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_WIDTH
        self.prs.slide_height = SLIDE_HEIGHT

    def slide(self, notes: str = ""):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[_BLANK_LAYOUT])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = WHITE
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return slide

    @staticmethod
    def text(slide, left, top, width, height, paragraphs, size: int, color=TEXT, bold=False, name="Testo"):
        """*paragraphs*: stringhe o liste di run ``(testo, grassetto)``."""
        box = slide.shapes.add_textbox(left, top, width, height)
        box.name = name
        frame = box.text_frame
        frame.word_wrap = True
        for i, paragraph in enumerate(paragraphs):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.space_after = Pt(6)
            runs = paragraph if isinstance(paragraph, list) else [(paragraph, bold)]
            for content, is_bold in runs:
                run = p.add_run()
                run.text = content
                run.font.size = Pt(size)
                run.font.bold = is_bold
                run.font.color.rgb = color
                run.font.name = FONT
        return box

    def title(self, slide, text: str, size: int = 32):
        """Titolo su una sola riga (circa 66 caratteri a 28 pt): il resto
        dell'area resta all'immagine o al testo."""
        limit = int(1850 / size)
        return self.text(
            slide, MARGIN, Inches(0.45), SLIDE_WIDTH - 2 * MARGIN, Inches(0.9), [_truncate(text, limit)], size, bold=True, name="Titolo"
        )

    def picture(self, slide, path: Path | None) -> bool:
        """Immagine adattata all'area contenuti senza deformarla; False se
        manca o non è leggibile (la slide mostrerà il testo)."""
        if path is None or not Path(path).exists():
            return False
        try:
            from PIL import Image

            with Image.open(path) as image:
                width_px, height_px = image.size
        except Exception:
            return False
        scale = min(IMAGE_WIDTH / width_px, IMAGE_HEIGHT / height_px)
        width, height = int(width_px * scale), int(height_px * scale)
        left = int((SLIDE_WIDTH - width) / 2)
        picture = slide.shapes.add_picture(str(path), left, CONTENT_TOP, width, height)
        picture.name = "Immagine"
        return True

    def to_bytes(self) -> bytes:
        buffer = io.BytesIO()
        self.prs.save(buffer)
        return buffer.getvalue()


def _source_notes(step: dict, s: dict[str, str]) -> list[str]:
    """Righe "Fonte" delle note: file, tipo, data e timestamp per i video,
    numeri di slide; per le run senza consolidamento i soli timestamp."""
    lines = []
    for source in step.get("sources", []) or []:
        kind = s.get(f"kind_{source.get('kind')}", source.get("kind") or "")
        date = f"{s['report_recorded']} {source['recorded_at']}" if source.get("recorded_at") else s["report_no_date"]
        timestamps = ", ".join(source.get("timestamps") or [])
        lines.append(f"{s['pptx_source']}: {source.get('filename', '')} ({kind}, {date})" + (f" · {timestamps}" if timestamps else ""))
    if step.get("source_slides"):
        lines.append(f"{s['pptx_slides']}: {', '.join(str(n) for n in step['source_slides'])}")
    if not lines:
        timestamps = step.get("source_timestamps") or ([step["timestamp"]] if step.get("timestamp") else [])
        if timestamps:
            lines.append(f"{s['pptx_timestamps']}: {', '.join(timestamps)}")
    return lines


def build_pptx(
    data: dict,
    screenshot_map: dict[tuple[str, int], Path | None],
    concept_image_map: dict[str, Path] | None,
    language: str = "it",
) -> bytes:
    s = get_strings(language)
    deck = _Deck()
    concept_image_map = concept_image_map or {}
    concepts = data.get("concepts", []) or []
    procedures = data.get("procedures", []) or []
    body_width = SLIDE_WIDTH - 2 * MARGIN
    body_height = SLIDE_HEIGHT - CONTENT_TOP - MARGIN

    # Copertina
    cover = deck.slide(notes=data.get("summary", ""))
    deck.text(cover, MARGIN, Inches(2.0), body_width, Inches(0.5), [s["pptx_tutorial"].upper()], 16, ACCENT, True, "Etichetta")
    deck.text(cover, MARGIN, Inches(2.6), body_width, Inches(1.8), [_truncate(data.get("title", ""), 110)], 40, bold=True, name="Titolo")
    deck.text(cover, MARGIN, Inches(4.6), body_width, Inches(2.2), [_truncate(data.get("summary", ""), 360)], 18, SECONDARY, name="Sommario")

    # Agenda
    has_both = bool(concepts) and bool(procedures)
    agenda = [
        (f"[{s['html_concept_tag']}] " if has_both else "") + con.get("title", "") for con in concepts
    ] + [(f"[{s['html_procedure_tag']}] " if has_both else "") + proc.get("title", "") for proc in procedures]
    pages = _chunks(list(enumerate(agenda, 1)), AGENDA_PER_SLIDE)
    for page_index, page in enumerate(pages, 1):
        slide = deck.slide(notes="\n".join(f"{i}. {item}" for i, item in page))
        deck.title(slide, _paged(s["pptx_agenda"], page_index, len(pages)))
        deck.text(slide, MARGIN, CONTENT_TOP, body_width, body_height, [f"{i}. {_truncate(item, 100)}" for i, item in page], 18)

    # Prerequisiti
    prerequisites = data.get("prerequisites", []) or []
    pages = _chunks(prerequisites, GLOSSARY_PER_SLIDE)
    for page_index, page in enumerate(pages, 1):
        slide = deck.slide(notes="\n".join(f"- {p}" for p in page))
        deck.title(slide, _paged(s["pptx_prerequisites"], page_index, len(pages)))
        deck.text(slide, MARGIN, CONTENT_TOP, body_width, body_height, [f"• {_truncate(p, 200)}" for p in page], 18)

    # Concetti: immagine della slide se c'è, altrimenti il testo
    for con in concepts:
        slide = deck.slide(notes=f"{con.get('title', '')}\n\n{con.get('content', '')}")
        deck.title(slide, con.get("title", ""))
        if not deck.picture(slide, concept_image_map.get(con.get("id"))):
            deck.text(slide, MARGIN, CONTENT_TOP, body_width, body_height, [_truncate(con.get("content", ""), 1100)], 18)

    # Procedure: slide di sezione, poi una slide per step
    for proc_index, proc in enumerate(procedures, 1):
        steps = proc.get("steps", []) or []
        section = deck.slide(notes=proc.get("title", ""))
        deck.text(
            section, MARGIN, Inches(2.3), body_width, Inches(0.5),
            [s["pptx_procedure_of"].format(i=proc_index, n=len(procedures)).upper()], 16, ACCENT, True, "Etichetta",
        )
        deck.text(section, MARGIN, Inches(2.9), body_width, Inches(1.8), [_truncate(proc.get("title", ""), 110)], 40, bold=True, name="Titolo")
        deck.text(section, MARGIN, Inches(4.8), body_width, Inches(0.6), [s["pptx_steps_count"].format(n=len(steps))], 20, SECONDARY, name="Step")

        for step in steps:
            description = step.get("description", "")
            notes = [step.get("title", ""), f"{s['pptx_description']}: {description}"]
            if step.get("notes"):
                notes.append(f"{s['pptx_step_notes']}: {step['notes']}")
            notes.extend(_source_notes(step, s))
            slide = deck.slide(notes="\n\n".join(notes))
            deck.title(slide, s["pptx_step_title"].format(k=step.get("number"), n=len(steps), title=step.get("title", "")), 28)
            if not deck.picture(slide, screenshot_map.get((proc.get("id"), step.get("number")))):
                paragraphs = [_truncate(description, 1100)]
                if step.get("notes"):
                    paragraphs.append(_truncate(step["notes"], 300))
                deck.text(slide, MARGIN, CONTENT_TOP, body_width, body_height, paragraphs, 18)

    # Glossario
    glossary = data.get("glossary", []) or []
    pages = _chunks(glossary, GLOSSARY_PER_SLIDE)
    for page_index, page in enumerate(pages, 1):
        slide = deck.slide(notes="\n".join(f"{g.get('term', '')}: {g.get('definition', '')}" for g in page))
        deck.title(slide, _paged(s["pptx_glossary"], page_index, len(pages)))
        deck.text(
            slide, MARGIN, CONTENT_TOP, body_width, body_height,
            [[(f"{_truncate(g.get('term', ''), 60)}: ", True), (_truncate(g.get("definition", ""), 150), False)] for g in page],
            15,
        )

    # Domande e risposte e note generali
    notes = data.get("notes", []) or []
    pages = _chunks(notes, NOTES_PER_SLIDE)
    for page_index, page in enumerate(pages, 1):
        slide = deck.slide(notes="\n\n".join(str(n) for n in page))
        deck.title(slide, _paged(s["pptx_notes"], page_index, len(pages)))
        deck.text(slide, MARGIN, CONTENT_TOP, body_width, body_height, [f"• {_truncate(n, 220)}" for n in page], 15)

    return deck.to_bytes()
