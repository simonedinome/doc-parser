"""Estrazione testo locale da .xlsx (openpyxl): tutte le sheet, intestazione per sheet, cap righe."""

from __future__ import annotations

import io

from openpyxl import load_workbook


def extract_xlsx(data: bytes, max_rows_per_sheet: int) -> str:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sections: list[str] = []
    for sheet in workbook.worksheets:
        lines = [f"## {sheet.title}"]
        truncated = False
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= max_rows_per_sheet:
                truncated = True
                break
            values = ["" if v is None else str(v) for v in row]
            lines.append(" | ".join(values))
        if truncated:
            lines.append(f"[Troncato a {max_rows_per_sheet} righe]")
        sections.append("\n".join(lines))
    return "\n\n".join(sections)
