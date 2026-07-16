"""Estrazione testo locale da .docx (python-docx): niente invio di binari Office ai modelli."""

from __future__ import annotations

import io

from docx import Document


def extract_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    for table in document.tables:
        lines.append("")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            lines.append(" | ".join(cells))
    return "\n".join(lines)
