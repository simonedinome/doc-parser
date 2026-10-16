import io
import json
import zipfile
from pathlib import Path
from unittest.mock import patch

import pytest

from doc_parser.core.models import RunRecord, RunStatus
from doc_parser.pipeline.core import process_run
from doc_parser.pipeline.step2_validate import Step2ValidationError
from doc_parser.providers.base import BaseProvider
from doc_parser.settings import Settings
from doc_parser.storage.sqlite_repo import SqliteRepository

_MODELS_FILE = Path(__file__).resolve().parent.parent.parent / "models.yaml"
STEP1_MODEL = "gemini-3.1-pro-preview"  # step1_video_model in models.yaml
OPENROUTER_MODEL = "anthropic/claude-haiku-4.5"
GEMINI_LITE_MODEL = "gemini-2.5-flash-lite"


class ScriptedProvider(BaseProvider):
    name = "scripted"

    def __init__(self, responses: list[tuple[str, int, int]]) -> None:
        self._responses = list(responses)
        self.calls: list[tuple[str, str, list]] = []

    def generate(self, system_prompt, user_prompt, attachments=None):
        self.calls.append((system_prompt, user_prompt, attachments or []))
        if not self._responses:
            raise AssertionError("Nessuna risposta scriptata rimasta")
        return self._responses.pop(0)


def _patch_providers(monkeypatch, by_model: dict[str, ScriptedProvider]) -> None:
    def _fake_get_provider_for_model(model_id, settings):
        return by_model[model_id]

    monkeypatch.setattr("doc_parser.pipeline.core.get_provider_for_model", _fake_get_provider_for_model)


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return Settings(_env_file=None, data_dir=tmp_path / "data", models_file=_MODELS_FILE, retry_max_attempts=2)


@pytest.fixture
def repo(settings: Settings) -> SqliteRepository:
    return SqliteRepository(data_dir=settings.data_dir)


def _flat_steps_response(steps: list[dict], title: str = "Titolo") -> str:
    return json.dumps(
        {
            "title": title,
            "summary": "Sommario",
            "prerequisites": [],
            "steps": steps,
            "glossary": [],
            "notes": [],
        }
    )


def _video_run(
    repo: SqliteRepository,
    model_id: str,
    content_type: str = "screencast",
    language: str = "it",
    output_artifacts: str = "html_standalone,procedura,glossario,rag_chunks",
) -> RunRecord:
    run = RunRecord(
        session_id="s1",
        requester_upn="utente@demo.local",
        input_filename="video.mp4",
        input_type="video",
        content_type=content_type,
        model_id=model_id,
        language=language,
        output_artifacts=output_artifacts,
    )
    run.input_path = repo.save_input_file(run.run_id, "video.mp4", b"video-bytes-fittizi")
    return repo.create_run(run)


def test_video_two_step_rewrite(monkeypatch, repo, settings) -> None:
    run = _video_run(repo, OPENROUTER_MODEL)
    steps = [{"number": 1, "timestamp": "00:10", "title": "S1", "description": "d1", "notes": "", "embedding_keywords": []}]
    step1 = ScriptedProvider([(_flat_steps_response(steps), 1000, 200)])
    # Lo step 2 riceve e restituisce il JSON già normalizzato a procedures[] (non lo
    # schema flat del profilo originale): stesso "id"/"number"/"timestamp", prosa riscritta.
    rewritten = {
        "title": "Titolo migliorato",
        "summary": "Sommario migliorato",
        "prerequisites": [],
        "procedures": [
            {
                "id": "proc_001",
                "title": "Titolo migliorato",
                "steps": [
                    {"number": 1, "timestamp": "00:10", "title": "S1 migliorato", "description": "d1 meglio", "notes": "", "embedding_keywords": []}
                ],
            }
        ],
        "glossary": [],
        "notes": [],
    }
    step2 = ScriptedProvider([(json.dumps(rewritten), 500, 300)])
    _patch_providers(monkeypatch, {STEP1_MODEL: step1, OPENROUTER_MODEL: step2})

    process_run(run, repo, settings)

    updated = repo.get_run(run.run_id)
    assert updated.status == RunStatus.DONE
    assert updated.raw_json_ref and updated.final_json_ref
    assert updated.raw_json_ref != updated.final_json_ref
    raw = json.loads(repo.get_file(updated.raw_json_ref))
    final = json.loads(repo.get_file(updated.final_json_ref))
    assert raw["title"] == "Titolo"
    assert final["title"] == "Titolo migliorato"
    assert updated.step1_model == STEP1_MODEL
    assert updated.step2_model == OPENROUTER_MODEL
    assert updated.step1_tokens_in == 1000 and updated.step2_tokens_in == 500
    assert updated.tokens_in == 1500 and updated.tokens_out == 500
    assert len(step1.calls) == 1
    assert len(step2.calls) == 1

    # Costo aggregato tra due provider diversi (Gemini step1 + Anthropic/OpenRouter
    # step2): verifica la matematica esatta, non solo "> 0".
    from doc_parser.core.pricing import compute_cost

    step1_expected = compute_cost(STEP1_MODEL, 1000, 200, settings.models_file)
    step2_expected = compute_cost(OPENROUTER_MODEL, 500, 300, settings.models_file)
    assert updated.step1_cost_eur == pytest.approx(step1_expected.cost_eur)
    assert updated.step2_cost_eur == pytest.approx(step2_expected.cost_eur)
    assert updated.cost_eur == pytest.approx(step1_expected.cost_eur + step2_expected.cost_eur)
    assert updated.cost_eur > 0

    assert updated.zip_ref and updated.output_dir_ref
    assert updated.output_path == updated.zip_ref
    assert updated.output_size_bytes and updated.output_size_bytes > 0
    zip_bytes = repo.get_file(updated.zip_ref)
    names = zipfile.ZipFile(io.BytesIO(zip_bytes)).namelist()
    assert {"procedura.txt", "glossario.txt", "rag_chunks.jsonl", "documentazione_standalone.html", "log.txt"} <= set(names)
    artifacts_dir = Path(updated.output_dir_ref)
    assert (artifacts_dir / "procedura.txt").exists()
    assert "Titolo migliorato" in (artifacts_dir / "procedura.txt").read_text(encoding="utf-8")


def test_video_folder_html_mode_includes_screenshots_and_english_labels(monkeypatch, repo, settings) -> None:
    run = _video_run(repo, GEMINI_LITE_MODEL, language="en", output_artifacts="html_folder,procedura,glossario,rag_chunks")
    steps = [{"number": 1, "timestamp": "00:01", "title": "S1", "description": "d1", "notes": "", "embedding_keywords": []}]
    provider = ScriptedProvider([(_flat_steps_response(steps), 100, 10)])
    _patch_providers(monkeypatch, {GEMINI_LITE_MODEL: provider})

    def _fake_ffmpeg_run(cmd, **kwargs):
        import subprocess

        if cmd[0] == "ffprobe":
            return subprocess.CompletedProcess(cmd, 0, stdout="120.0", stderr="")
        Path(cmd[-1]).write_bytes(b"PNG")
        return subprocess.CompletedProcess(cmd, 0)

    with patch("doc_parser.pipeline.screenshots.subprocess.run", side_effect=_fake_ffmpeg_run):
        process_run(run, repo, settings)

    updated = repo.get_run(run.run_id)
    assert updated.status == RunStatus.DONE
    zip_bytes = repo.get_file(updated.zip_ref)
    zf = zipfile.ZipFile(io.BytesIO(zip_bytes))
    names = zf.namelist()
    assert "screenshots/proc_001_step_001.png" in names
    assert "documentazione.html" in names
    assert "documentazione_standalone.html" not in names
    procedura = zf.read("procedura.txt").decode("utf-8")
    assert "SUMMARY" in procedura and "PROCEDURE 1" in procedura


def test_video_single_pass_gemini_skips_step2(monkeypatch, repo, settings) -> None:
    run = _video_run(repo, GEMINI_LITE_MODEL)
    steps = [{"number": 1, "timestamp": "00:10", "title": "S1", "description": "d1", "notes": "", "embedding_keywords": []}]
    chosen_gemini = ScriptedProvider([(_flat_steps_response(steps), 800, 150)])
    fixed_step1_default = ScriptedProvider([])  # non deve mai essere chiamato
    _patch_providers(monkeypatch, {STEP1_MODEL: fixed_step1_default, GEMINI_LITE_MODEL: chosen_gemini})

    process_run(run, repo, settings)

    updated = repo.get_run(run.run_id)
    assert updated.status == RunStatus.DONE
    assert updated.raw_json_ref == updated.final_json_ref
    assert updated.step1_model == GEMINI_LITE_MODEL
    assert updated.step2_model is None
    assert len(chosen_gemini.calls) == 1
    assert len(fixed_step1_default.calls) == 0


def test_video_multi_file_merge(monkeypatch, repo, settings) -> None:
    run = RunRecord(
        session_id="s1", requester_upn="utente@demo.local", input_filename="run",
        input_type="video", content_type="screencast", model_id=GEMINI_LITE_MODEL,
    )
    ref_a = repo.save_input_file(run.run_id, "a.mp4", b"video-a")
    ref_b = repo.save_input_file(run.run_id, "b.mp4", b"video-b")
    manifest = [{"filename": "a.mp4", "path_ref": ref_a}, {"filename": "b.mp4", "path_ref": ref_b}]
    run.input_manifest_ref = repo.save_input_file(
        run.run_id, "manifest.json", json.dumps(manifest).encode("utf-8")
    )
    repo.create_run(run)

    steps_a = [{"number": 1, "timestamp": "00:10", "title": "A1", "description": "d", "notes": "", "embedding_keywords": []}]
    steps_b = [{"number": 1, "timestamp": "00:05", "title": "B1", "description": "d", "notes": "", "embedding_keywords": []}]
    provider = ScriptedProvider(
        [(_flat_steps_response(steps_a, title="Video A"), 100, 10), (_flat_steps_response(steps_b, title="Video B"), 100, 10)]
    )
    _patch_providers(monkeypatch, {GEMINI_LITE_MODEL: provider})

    with patch("doc_parser.pipeline.core.screenshots.extract_screenshots", return_value={}):
        process_run(run, repo, settings)

    updated = repo.get_run(run.run_id)
    assert updated.status == RunStatus.DONE
    merged = json.loads(repo.get_file(updated.final_json_ref))
    assert [p["id"] for p in merged["procedures"]] == ["proc_001", "proc_002"]
    assert len(provider.calls) == 2


def test_step2_validation_failure_marks_failed_and_keeps_raw(monkeypatch, repo, settings) -> None:
    run = _video_run(repo, OPENROUTER_MODEL)
    steps = [{"number": 1, "timestamp": "00:10", "title": "S1", "description": "d1", "notes": "", "embedding_keywords": []}]
    step1 = ScriptedProvider([(_flat_steps_response(steps), 100, 10)])
    bad_steps = [{"number": 99, "timestamp": "99:99", "title": "Violazione", "description": "d", "notes": "", "embedding_keywords": []}]
    bad_response = _flat_steps_response(bad_steps)
    step2 = ScriptedProvider([(bad_response, 10, 10)] * (settings.retry_max_attempts + 1))
    _patch_providers(monkeypatch, {STEP1_MODEL: step1, OPENROUTER_MODEL: step2})

    with pytest.raises(Step2ValidationError):
        process_run(run, repo, settings)

    updated = repo.get_run(run.run_id)
    assert updated.raw_json_ref is not None  # il grezzo resta disponibile per Rielabora
    assert updated.completed_at is None  # finalizzato in FAILED dal worker (non testato qui), non da process_run
    assert len(step2.calls) == settings.retry_max_attempts + 1


def _document_run(
    repo: SqliteRepository,
    filename: str,
    data: bytes,
    input_type: str,
    model_id: str,
    output_mode: str = "documentazione",
    output_artifacts: str = "html_standalone,procedura,glossario,rag_chunks",
    content_type: str | None = None,
) -> RunRecord:
    run = RunRecord(
        session_id="s1", requester_upn="utente@demo.local", input_filename=filename,
        input_type=input_type, model_id=model_id, output_mode=output_mode,
        output_artifacts=output_artifacts, content_type=content_type,
    )
    run.input_path = repo.save_input_file(run.run_id, filename, data)
    return repo.create_run(run)


def _sample_pdf_bytes(pages: int = 1) -> bytes:
    import fitz

    doc = fitz.open()
    for i in range(pages):
        doc.new_page().insert_text((72, 72), f"Pagina {i + 1}")
    data = doc.tobytes()
    doc.close()
    return data


def test_document_pdf_single_step(monkeypatch, repo, settings) -> None:
    run = _document_run(repo, "manuale.pdf", _sample_pdf_bytes(), "pdf", OPENROUTER_MODEL)
    steps = [{"number": 1, "timestamp": "Pagina 1", "title": "S1", "description": "d", "notes": "", "embedding_keywords": []}]
    provider = ScriptedProvider([(_flat_steps_response(steps), 400, 100)])
    _patch_providers(monkeypatch, {OPENROUTER_MODEL: provider})

    process_run(run, repo, settings)

    updated = repo.get_run(run.run_id)
    assert updated.status == RunStatus.DONE
    assert updated.raw_json_ref == updated.final_json_ref
    assert updated.input_pages == 1
    assert updated.screenshots_ref is not None
    assert len(provider.calls[0][2]) == 1  # un allegato PDF nativo
    assert provider.calls[0][2][0].mime_type == "application/pdf"

    assert updated.zip_ref is not None
    zip_bytes = repo.get_file(updated.zip_ref)
    zf = zipfile.ZipFile(io.BytesIO(zip_bytes))
    # html_mode di default e' "standalone": l'immagine di pagina 1 e' embedded in base64,
    # non in una cartella screenshots/ separata.
    html = zf.read("documentazione_standalone.html").decode("utf-8")
    assert "data:image/png;base64," in html


def test_document_docx_uses_local_extraction_no_attachment(monkeypatch, repo, settings) -> None:
    import io

    from docx import Document

    document = Document()
    document.add_paragraph("Contenuto del docx di test")
    buffer = io.BytesIO()
    document.save(buffer)

    run = _document_run(repo, "relazione.docx", buffer.getvalue(), "docx", OPENROUTER_MODEL)
    steps = [{"number": 1, "timestamp": "Pagina 1", "title": "S1", "description": "d", "notes": "", "embedding_keywords": []}]
    provider = ScriptedProvider([(_flat_steps_response(steps), 200, 50)])
    _patch_providers(monkeypatch, {OPENROUTER_MODEL: provider})

    process_run(run, repo, settings)

    updated = repo.get_run(run.run_id)
    assert updated.status == RunStatus.DONE
    system_prompt, user_prompt, attachments = provider.calls[0]
    assert attachments == []
    assert "Contenuto del docx di test" in user_prompt


def test_images_multiple_become_single_document(monkeypatch, repo, settings) -> None:
    run = RunRecord(
        session_id="s1", requester_upn="utente@demo.local", input_filename="run-immagini",
        input_type="image", model_id=OPENROUTER_MODEL,
    )
    ref_a = repo.save_input_file(run.run_id, "a.png", b"\x89PNG-a")
    ref_b = repo.save_input_file(run.run_id, "b.png", b"\x89PNG-b")
    manifest = [{"filename": "a.png", "path_ref": ref_a}, {"filename": "b.png", "path_ref": ref_b}]
    run.input_manifest_ref = repo.save_input_file(run.run_id, "manifest.json", json.dumps(manifest).encode("utf-8"))
    repo.create_run(run)

    steps = [{"number": 1, "timestamp": "Pagina 1", "title": "S1", "description": "d", "notes": "", "embedding_keywords": []}]
    provider = ScriptedProvider([(_flat_steps_response(steps), 300, 80)])
    _patch_providers(monkeypatch, {OPENROUTER_MODEL: provider})

    process_run(run, repo, settings)

    updated = repo.get_run(run.run_id)
    assert updated.status == RunStatus.DONE
    assert len(provider.calls[0][2]) == 2


def test_document_multi_file_group_merge(monkeypatch, repo, settings) -> None:
    import io

    from docx import Document

    document = Document()
    document.add_paragraph("Contenuto docx")
    docx_buffer = io.BytesIO()
    document.save(docx_buffer)

    run = RunRecord(
        session_id="s1", requester_upn="utente@demo.local", input_filename="run-documenti",
        input_type="document", model_id=OPENROUTER_MODEL,
    )
    ref_pdf = repo.save_input_file(run.run_id, "a.pdf", _sample_pdf_bytes())
    ref_docx = repo.save_input_file(run.run_id, "b.docx", docx_buffer.getvalue())
    manifest = [{"filename": "a.pdf", "path_ref": ref_pdf}, {"filename": "b.docx", "path_ref": ref_docx}]
    run.input_manifest_ref = repo.save_input_file(run.run_id, "manifest.json", json.dumps(manifest).encode("utf-8"))
    repo.create_run(run)

    steps_a = [{"number": 1, "timestamp": "Pagina 1", "title": "A1", "description": "d", "notes": "", "embedding_keywords": []}]
    steps_b = [{"number": 1, "timestamp": "Pagina 1", "title": "B1", "description": "d", "notes": "", "embedding_keywords": []}]
    provider = ScriptedProvider(
        [(_flat_steps_response(steps_a, title="Doc A"), 100, 10), (_flat_steps_response(steps_b, title="Doc B"), 100, 10)]
    )
    _patch_providers(monkeypatch, {OPENROUTER_MODEL: provider})

    process_run(run, repo, settings)

    updated = repo.get_run(run.run_id)
    assert updated.status == RunStatus.DONE
    merged = json.loads(repo.get_file(updated.final_json_ref))
    assert [p["id"] for p in merged["procedures"]] == ["proc_001", "proc_002"]
    assert len(provider.calls) == 2  # una chiamata LLM per file, poi merge
    assert updated.screenshots_ref is not None  # screenshot estratti solo per il PDF del gruppo


def test_document_output_artifacts_selection_filters_zip_contents(monkeypatch, repo, settings) -> None:
    run = _document_run(
        repo, "manuale.pdf", _sample_pdf_bytes(), "pdf", OPENROUTER_MODEL,
        output_artifacts="procedura,rag_chunks",
    )
    steps = [{"number": 1, "timestamp": "Pagina 1", "title": "S1", "description": "d", "notes": "", "embedding_keywords": []}]
    provider = ScriptedProvider([(_flat_steps_response(steps), 400, 100)])
    _patch_providers(monkeypatch, {OPENROUTER_MODEL: provider})

    process_run(run, repo, settings)

    updated = repo.get_run(run.run_id)
    assert updated.status == RunStatus.DONE
    zip_bytes = repo.get_file(updated.zip_ref)
    names = set(zipfile.ZipFile(io.BytesIO(zip_bytes)).namelist())
    assert names == {"procedura.txt", "rag_chunks.jsonl", "log.txt"}


def test_manuale_profile_rejects_non_pdf_input(monkeypatch, repo, settings) -> None:
    from doc_parser.core.errors import GuardrailRejection

    import io

    from docx import Document

    document = Document()
    document.add_paragraph("Contenuto docx")
    buffer = io.BytesIO()
    document.save(buffer)

    run = _document_run(
        repo, "manuale.docx", buffer.getvalue(), "docx", OPENROUTER_MODEL, content_type="manuale"
    )
    _patch_providers(monkeypatch, {})  # non deve mai essere chiamato

    with pytest.raises(GuardrailRejection, match="'manuale' accetta solo PDF"):
        process_run(run, repo, settings)


def test_documento_and_manuale_runs_on_same_pdf_are_not_deduped(monkeypatch, repo, settings) -> None:
    pdf_bytes = _sample_pdf_bytes()
    flat_steps = [{"number": 1, "timestamp": "Pagina 1", "title": "S1", "description": "d", "notes": "", "embedding_keywords": []}]
    manuale_response = json.dumps(
        {
            "title": "T", "summary": "S", "prerequisites": [],
            "concepts": [{"id": "con_001", "title": "C1", "content": "...", "embedding_keywords": []}],
            "procedures": [],
            "glossary": [], "notes": [],
        }
    )
    provider = ScriptedProvider([(_flat_steps_response(flat_steps), 100, 20), (manuale_response, 100, 20)])
    _patch_providers(monkeypatch, {OPENROUTER_MODEL: provider})

    first = _document_run(repo, "manuale.pdf", pdf_bytes, "pdf", OPENROUTER_MODEL, content_type="documento")
    process_run(first, repo, settings)
    first = repo.get_run(first.run_id)
    assert first.status == RunStatus.DONE

    second = _document_run(repo, "manuale.pdf", pdf_bytes, "pdf", OPENROUTER_MODEL, content_type="manuale")
    process_run(second, repo, settings)
    second = repo.get_run(second.run_id)

    assert second.status == RunStatus.DONE  # non DUPLICATE: content_type diverso
    assert len(provider.calls) == 2  # entrambe le run hanno chiamato l'LLM


def test_kb_densa_output_mode(monkeypatch, repo, settings) -> None:
    run = _document_run(repo, "manuale.pdf", _sample_pdf_bytes(), "pdf", OPENROUTER_MODEL, output_mode="kb_densa")
    chunks_response = json.dumps(
        {"chunks": [{"index": 1, "title": "Sezione 1", "text": "Testo denso.", "has_uncertainty": False}]}
    )
    provider = ScriptedProvider([(chunks_response, 300, 60)])
    _patch_providers(monkeypatch, {OPENROUTER_MODEL: provider})

    process_run(run, repo, settings)

    updated = repo.get_run(run.run_id)
    assert updated.status == RunStatus.DONE
    assert updated.output_path is not None
    assert updated.chunk_count == 1
    assert updated.uncertainty_count == 0
    txt = repo.get_file(updated.output_path).decode("utf-8")
    assert "Testo denso." in txt


def test_elaborate_only_reuses_parent_raw_json(monkeypatch, repo, settings) -> None:
    parent = _video_run(repo, GEMINI_LITE_MODEL)
    raw_data = {
        "title": "Titolo grezzo",
        "procedures": [{"id": "proc_001", "title": "P1", "steps": [{"number": 1, "timestamp": "00:10", "title": "S1", "description": "d"}]}],
    }
    raw_ref = repo.save_output_file(parent.run_id, "raw.json", json.dumps(raw_data).encode("utf-8"))
    repo.update_run(parent.run_id, raw_json_ref=raw_ref, screenshots_ref="/tmp/screens/parent", status=RunStatus.DONE)

    child = RunRecord(
        session_id="s1", requester_upn="utente@demo.local", input_filename="",
        input_type="video", model_id=OPENROUTER_MODEL, parent_run_id=parent.run_id,
    )
    repo.create_run(child)

    rewritten = {
        "title": "Titolo rifinito",
        "procedures": [{"id": "proc_001", "title": "P1", "steps": [{"number": 1, "timestamp": "00:10", "title": "S1 meglio", "description": "d meglio"}]}],
    }
    step2 = ScriptedProvider([(json.dumps(rewritten), 200, 100)])
    _patch_providers(monkeypatch, {OPENROUTER_MODEL: step2})

    process_run(child, repo, settings)

    updated = repo.get_run(child.run_id)
    assert updated.status == RunStatus.DONE
    assert updated.raw_json_ref == raw_ref
    assert updated.screenshots_ref == "/tmp/screens/parent"
    final = json.loads(repo.get_file(updated.final_json_ref))
    assert final["title"] == "Titolo rifinito"
    assert updated.zip_ref is not None
    zip_bytes = repo.get_file(updated.zip_ref)
    procedura = zipfile.ZipFile(io.BytesIO(zip_bytes)).read("procedura.txt").decode("utf-8")
    assert "Titolo rifinito" in procedura


def test_elaborate_only_missing_parent_raises_guardrail(repo, settings) -> None:
    from doc_parser.core.errors import GuardrailRejection

    child = RunRecord(
        session_id="s1", requester_upn="utente@demo.local", input_filename="",
        input_type="video", model_id=OPENROUTER_MODEL, parent_run_id="run-inesistente",
    )
    repo.create_run(child)

    with pytest.raises(GuardrailRejection):
        process_run(child, repo, settings)


def test_elaborate_only_missing_artifact_on_disk_raises_guardrail(repo, settings) -> None:
    from doc_parser.core.errors import GuardrailRejection

    parent = _video_run(repo, GEMINI_LITE_MODEL)
    repo.update_run(parent.run_id, raw_json_ref=str(settings.data_dir / "outputs" / "assente.json"), status=RunStatus.DONE)

    child = RunRecord(
        session_id="s1", requester_upn="utente@demo.local", input_filename="",
        input_type="video", model_id=OPENROUTER_MODEL, parent_run_id=parent.run_id,
    )
    repo.create_run(child)

    with pytest.raises(GuardrailRejection):
        process_run(child, repo, settings)


def test_missing_input_raises_guardrail(repo, settings) -> None:
    from doc_parser.core.errors import GuardrailRejection

    run = RunRecord(session_id="s1", requester_upn="u@x", input_filename="", input_type="pdf", model_id=OPENROUTER_MODEL)
    repo.create_run(run)

    with pytest.raises(GuardrailRejection, match="mancante"):
        process_run(run, repo, settings)


def test_dedup_reuses_output_for_identical_file_and_config(monkeypatch, repo, settings) -> None:
    pdf_bytes = _sample_pdf_bytes()
    steps = [{"number": 1, "timestamp": "Pagina 1", "title": "S1", "description": "d", "notes": "", "embedding_keywords": []}]
    provider = ScriptedProvider([(_flat_steps_response(steps), 100, 20)])
    _patch_providers(monkeypatch, {OPENROUTER_MODEL: provider})

    first = _document_run(repo, "manuale.pdf", pdf_bytes, "pdf", OPENROUTER_MODEL)
    process_run(first, repo, settings)
    first = repo.get_run(first.run_id)
    assert first.status == RunStatus.DONE

    second = _document_run(repo, "copia.pdf", pdf_bytes, "pdf", OPENROUTER_MODEL)
    process_run(second, repo, settings)
    second = repo.get_run(second.run_id)

    assert second.status == RunStatus.DUPLICATE
    assert second.final_json_ref == first.final_json_ref
    assert second.zip_ref == first.zip_ref
    assert second.tokens_total == 0 and second.cost_eur == 0.0
    assert first.run_id in (second.error_message or "")
    assert len(provider.calls) == 1  # nessuna seconda chiamata LLM


def test_dedup_does_not_trigger_for_different_model(monkeypatch, repo, settings) -> None:
    pdf_bytes = _sample_pdf_bytes()
    steps = [{"number": 1, "timestamp": "Pagina 1", "title": "S1", "description": "d", "notes": "", "embedding_keywords": []}]
    provider_a = ScriptedProvider([(_flat_steps_response(steps), 100, 20)])
    provider_b = ScriptedProvider([(_flat_steps_response(steps), 100, 20)])
    _patch_providers(monkeypatch, {OPENROUTER_MODEL: provider_a, GEMINI_LITE_MODEL: provider_b})

    first = _document_run(repo, "manuale.pdf", pdf_bytes, "pdf", OPENROUTER_MODEL)
    process_run(first, repo, settings)

    second = _document_run(repo, "manuale.pdf", pdf_bytes, "pdf", GEMINI_LITE_MODEL)
    process_run(second, repo, settings)
    second = repo.get_run(second.run_id)

    assert second.status == RunStatus.DONE
    assert len(provider_a.calls) == 1
    assert len(provider_b.calls) == 1


def test_dedup_does_not_trigger_for_different_output_mode(monkeypatch, repo, settings) -> None:
    pdf_bytes = _sample_pdf_bytes()
    steps = [{"number": 1, "timestamp": "Pagina 1", "title": "S1", "description": "d", "notes": "", "embedding_keywords": []}]
    chunks_response = json.dumps({"chunks": [{"index": 1, "title": "T", "text": "Testo.", "has_uncertainty": False}]})
    provider = ScriptedProvider([(_flat_steps_response(steps), 100, 20), (chunks_response, 50, 10)])
    _patch_providers(monkeypatch, {OPENROUTER_MODEL: provider})

    first = _document_run(repo, "manuale.pdf", pdf_bytes, "pdf", OPENROUTER_MODEL, output_mode="documentazione")
    process_run(first, repo, settings)

    second = _document_run(repo, "manuale.pdf", pdf_bytes, "pdf", OPENROUTER_MODEL, output_mode="kb_densa")
    process_run(second, repo, settings)
    second = repo.get_run(second.run_id)

    assert second.status == RunStatus.DONE
    assert len(provider.calls) == 2


def test_step2_retry_resends_raw_json_with_feedback(monkeypatch, repo, settings) -> None:
    run = _video_run(repo, OPENROUTER_MODEL)
    steps = [{"number": 1, "timestamp": "00:10", "title": "S1", "description": "d1", "notes": "", "embedding_keywords": []}]
    step1 = ScriptedProvider([(_flat_steps_response(steps), 100, 10)])
    bad = _flat_steps_response([{"number": 99, "timestamp": "99:99", "title": "X", "description": "d"}])
    good = {
        "title": "Rifinito", "summary": "S", "prerequisites": [], "glossary": [], "notes": [],
        "procedures": [{"id": "proc_001", "title": "Titolo", "steps": [{"number": 1, "timestamp": "00:10", "title": "S1", "description": "d1 meglio"}]}],
    }
    step2 = ScriptedProvider([(bad, 10, 10), (json.dumps(good), 10, 10)])
    _patch_providers(monkeypatch, {STEP1_MODEL: step1, OPENROUTER_MODEL: step2})

    process_run(run, repo, settings)

    updated = repo.get_run(run.run_id)
    assert updated.status == RunStatus.DONE
    assert updated.retry_count == 1
    retry_prompt = step2.calls[1][1]
    assert retry_prompt.startswith("Ecco il JSON grezzo da rifinire:")
    assert '"timestamp": "00:10"' in retry_prompt  # il JSON grezzo viaggia di nuovo
    assert "ha violato il vincolo" in retry_prompt


def test_pptx_requested_for_document_run_is_skipped_with_warning(monkeypatch, repo, settings) -> None:
    run = _document_run(repo, "manuale.pdf", _sample_pdf_bytes(), "pdf", OPENROUTER_MODEL, output_artifacts="procedura,pptx")
    steps = [{"number": 1, "timestamp": "Pagina 1", "title": "S1", "description": "d", "notes": "", "embedding_keywords": []}]
    _patch_providers(monkeypatch, {OPENROUTER_MODEL: ScriptedProvider([(_flat_steps_response(steps), 10, 10)])})

    process_run(run, repo, settings)

    zf = zipfile.ZipFile(io.BytesIO(repo.get_file(repo.get_run(run.run_id).zip_ref)))
    assert "presentazione_tutorial.pptx" not in zf.namelist()
    assert "Presentazione PPT non generata" in zf.read("log.txt").decode("utf-8")


def test_single_video_run_produces_pptx_when_requested(monkeypatch, repo, settings) -> None:
    run = _video_run(repo, GEMINI_LITE_MODEL, content_type="video", output_artifacts="pptx")
    response = json.dumps(
        {"title": "T", "summary": "S", "prerequisites": [], "glossary": [], "notes": [],
         "procedures": [{"id": "proc_001", "title": "P", "steps": [{"number": 1, "title": "S1", "description": "d", "source_timestamps": ["00:10"]}]}]}
    )
    _patch_providers(monkeypatch, {GEMINI_LITE_MODEL: ScriptedProvider([(response, 10, 10)])})

    with patch("doc_parser.pipeline.screenshots.subprocess.run", side_effect=FileNotFoundError()):
        process_run(run, repo, settings)

    names = zipfile.ZipFile(io.BytesIO(repo.get_file(repo.get_run(run.run_id).zip_ref))).namelist()
    assert "presentazione_tutorial.pptx" in names
