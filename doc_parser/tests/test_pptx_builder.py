import io
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from doc_parser.output.pptx_builder import build_pptx


@pytest.fixture
def screenshot(tmp_path: Path) -> Path:
    from PIL import Image

    path = tmp_path / "proc_001_step_001.png"
    Image.new("RGB", (1920, 1080), (200, 210, 220)).save(path)
    return path


def _data() -> dict:
    return {
        "title": "Manuale tariffe",
        "summary": "Gestione delle tariffe.",
        "prerequisites": ["Accesso"],
        "concepts": [{"id": "con_001", "title": "Tariffe", "content": "Cosa sono le tariffe."}],
        "procedures": [
            {
                "id": "proc_001",
                "title": "Creare una tariffa",
                "steps": [
                    {"number": 1, "title": "Aprire Tariffe", "description": "Navigare in Impostazioni > Tariffe.", "notes": "",
                     "sources": [{"input_id": "v1", "filename": "webinar.mp4", "kind": "webinar", "recorded_at": "2026-06-03", "timestamps": ["00:10"]}],
                     "source_slides": [2]},
                    {"number": 2, "title": "Salvare", "description": "Cliccare su Salva.", "notes": "Controllare le date.",
                     "source_timestamps": ["01:00"]},
                ],
            },
            {"id": "proc_002", "title": "Verificare", "steps": [{"number": 1, "title": "Aprire Calendario", "description": "Aprire il calendario."}]},
        ],
        "glossary": [{"term": f"T{i}", "definition": "D"} for i in range(1, 10)],
        "notes": [f"Domanda {i}? / Risposta {i}." for i in range(1, 8)],
    }


def _open(data_bytes: bytes):
    return Presentation(io.BytesIO(data_bytes))


def _texts(slide) -> str:
    return " ".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)


def _notes(slide) -> str:
    return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""


def test_slide_count_follows_structure(screenshot) -> None:
    prs = _open(build_pptx(_data(), {("proc_001", 1): screenshot}, {}, "it"))
    # copertina + agenda + prerequisiti + 1 concetto + (sezione + 2 step) + (sezione + 1 step)
    # + glossario 9 termini (2 slide) + 7 note (2 slide)
    assert len(prs.slides) == 1 + 1 + 1 + 1 + 3 + 2 + 2 + 2
    assert prs.slide_width == Inches(13.333) and prs.slide_height == Inches(7.5)
    titles = [_texts(slide) for slide in prs.slides]
    assert "PROCEDURA 1 DI 2" in titles[4] and "2 step" in titles[4]
    assert titles[9].startswith("Glossario (1/2)")
    assert titles[11].startswith("Domande e risposte e note (1/2)")


def test_every_step_slide_has_notes_with_description_and_source(screenshot) -> None:
    prs = _open(build_pptx(_data(), {("proc_001", 1): screenshot}, {}, "it"))
    step_slides = [s for s in prs.slides if _texts(s).startswith("Step ")]
    assert len(step_slides) == 3
    first, second, _ = (_notes(s) for s in step_slides)
    assert "Descrizione: Navigare in Impostazioni > Tariffe." in first
    assert "Fonte: webinar.mp4 (webinar, registrato il 2026-06-03) · 00:10" in first
    assert "Slide: 2" in first
    assert "Note: Controllare le date." in second and "Timestamp: 01:00" in second


def test_step_with_image_has_one_picture_fitted_without_distortion(screenshot) -> None:
    prs = _open(build_pptx(_data(), {("proc_001", 1): screenshot}, {}, "it"))
    slide = next(s for s in prs.slides if _texts(s).startswith("Step 1/2"))
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1
    picture = pictures[0]
    assert picture.height <= Inches(5.6) and picture.width <= Inches(12.3)
    assert abs(picture.width / picture.height - 1920 / 1080) < 0.01
    assert "Navigare in Impostazioni" not in _texts(slide)


def test_step_without_image_shows_text(screenshot) -> None:
    prs = _open(build_pptx(_data(), {("proc_001", 1): screenshot}, {}, "it"))
    slide = next(s for s in prs.slides if _texts(s).startswith("Step 2/2"))
    assert not any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
    assert "Cliccare su Salva." in _texts(slide)


def test_concept_uses_slide_image_when_available(tmp_path: Path, screenshot) -> None:
    prs = _open(build_pptx(_data(), {}, {"con_001": screenshot}, "it"))
    slide = next(s for s in prs.slides if _texts(s).startswith("Tariffe"))
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
    assert "Cosa sono le tariffe." in _notes(slide)


def test_long_text_is_truncated_on_slide_and_complete_in_notes() -> None:
    data = _data()
    long_text = "Parola " * 400
    data["procedures"][1]["steps"][0]["description"] = long_text
    prs = _open(build_pptx(data, {}, {}, "it"))
    slide = next(s for s in prs.slides if _texts(s).startswith("Step 1/1"))
    assert "…" in _texts(slide)
    assert long_text.strip() in _notes(slide)


def test_english_labels() -> None:
    prs = _open(build_pptx(_data(), {}, {}, "en"))
    texts = [_texts(s) for s in prs.slides]
    assert "Agenda" in texts[1] and "Prerequisites" in texts[2]
    assert any("PROCEDURE 1 OF 2" in t for t in texts)
    assert any(t.startswith("Glossary (1/2)") for t in texts)
    assert any(t.startswith("Questions, answers and notes") for t in texts)
    assert any("2 steps" in t for t in texts)
