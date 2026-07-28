import io
import subprocess
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document
from openpyxl import Workbook

from doc_parser.core.errors import GuardrailRejection
from doc_parser.ingestion.extract import extract_text
from doc_parser.ingestion.extractors.doc_legacy_extractor import extract_doc
from doc_parser.ingestion.extractors.docx_extractor import extract_docx
from doc_parser.ingestion.extractors.text_extractor import extract_csv, extract_txt
from doc_parser.ingestion.extractors.xlsx_extractor import extract_xlsx


def _sample_docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Titolo del documento")
    document.add_paragraph("Un paragrafo di prova.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Colonna A"
    table.rows[0].cells[1].text = "Colonna B"
    table.rows[1].cells[0].text = "1"
    table.rows[1].cells[1].text = "2"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _sample_xlsx_bytes(rows: int = 3) -> bytes:
    workbook = Workbook()
    sheet1 = workbook.active
    sheet1.title = "Foglio1"
    sheet1.append(["Nome", "Valore"])
    for i in range(rows):
        sheet1.append([f"riga{i}", i])
    sheet2 = workbook.create_sheet("Foglio2")
    sheet2.append(["X"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_extract_docx_includes_paragraphs_and_tables() -> None:
    text = extract_docx(_sample_docx_bytes())
    assert "Titolo del documento" in text
    assert "Un paragrafo di prova." in text
    assert "Colonna A | Colonna B" in text
    assert "1 | 2" in text


def test_extract_xlsx_covers_all_sheets_with_header() -> None:
    text = extract_xlsx(_sample_xlsx_bytes(), max_rows_per_sheet=100)
    assert "## Foglio1" in text
    assert "## Foglio2" in text
    assert "Nome | Valore" in text
    assert "riga0 | 0" in text


def test_extract_xlsx_truncates_rows_per_sheet() -> None:
    text = extract_xlsx(_sample_xlsx_bytes(rows=10), max_rows_per_sheet=3)
    assert "[Troncato a 3 righe]" in text
    assert "riga5 | 5" not in text


def test_extract_txt_decodes_utf8() -> None:
    assert extract_txt("ciao è ù".encode("utf-8")) == "ciao è ù"


def test_extract_csv_normalizes_rows() -> None:
    data = "a,b,c\n1,2,3\n".encode("utf-8")
    assert extract_csv(data) == "a | b | c\n1 | 2 | 3"


def test_extract_text_dispatch() -> None:
    assert extract_text("txt", b"ciao") == "ciao"
    assert "a | b" in extract_text("csv", b"a,b\n")
    assert "## Foglio1" in extract_text("xlsx", _sample_xlsx_bytes())
    assert "Titolo del documento" in extract_text("docx", _sample_docx_bytes())


# La conversione soffice --headless è mockata (non eseguita realmente): il
# sandbox di test non garantisce un ambiente grafico/headless funzionante per
# LibreOffice, e la funzione va comunque verificata a contratto (comando
# invocato, file letto dall'outdir, errori gestiti).


def _fake_soffice_run(cmd, **kwargs):
    outdir = Path(cmd[cmd.index("--outdir") + 1])
    (outdir / "input.txt").write_text("Titolo del documento\nUn paragrafo di prova.", encoding="utf-8")
    return subprocess.CompletedProcess(cmd, 0)


def test_extract_doc_converts_via_soffice_and_reads_output() -> None:
    with patch(
        "doc_parser.ingestion.extractors.doc_legacy_extractor.subprocess.run",
        side_effect=_fake_soffice_run,
    ):
        text = extract_doc(b"contenuto .doc fittizio")
    assert "Titolo del documento" in text
    assert "Un paragrafo di prova." in text


def test_extract_doc_missing_soffice_raises_clear_error() -> None:
    with patch(
        "doc_parser.ingestion.extractors.doc_legacy_extractor.subprocess.run",
        side_effect=FileNotFoundError(),
    ):
        with pytest.raises(GuardrailRejection, match="LibreOffice"):
            extract_doc(b"contenuto qualsiasi")


def test_extract_doc_conversion_failure_raises_guardrail() -> None:
    with patch(
        "doc_parser.ingestion.extractors.doc_legacy_extractor.subprocess.run",
        side_effect=subprocess.CalledProcessError(1, "soffice"),
    ):
        with pytest.raises(GuardrailRejection, match="Conversione"):
            extract_doc(b"contenuto qualsiasi")


def test_extract_doc_dispatch_via_extract_text() -> None:
    with patch(
        "doc_parser.ingestion.extractors.doc_legacy_extractor.subprocess.run",
        side_effect=_fake_soffice_run,
    ):
        assert "Titolo del documento" in extract_text("doc", b"contenuto .doc fittizio")
